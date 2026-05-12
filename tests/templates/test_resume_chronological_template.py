"""Validate the chronological resume template is ATS-safe."""
from pathlib import Path

from scripts.validators.ats_check import ats_check

TEMPLATE = Path(__file__).parent.parent.parent / "templates" / "resume" / "chronological.docx"


def test_template_exists():
    assert TEMPLATE.exists(), f"Template missing: {TEMPLATE}"


def test_template_passes_ats_check():
    result = ats_check(TEMPLATE)
    assert result.passed, f"Template failed ATS check: {[f.code for f in result.failures]}"


def test_template_contains_required_placeholders():
    from docx import Document
    doc = Document(str(TEMPLATE))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    for placeholder in ("{{CANDIDATE_NAME}}", "{{CANDIDATE_CONTACT_LINE}}",
                        "{{SUMMARY}}", "{{SKILLS}}", "{{EXPERIENCE}}", "{{EDUCATION}}"):
        assert placeholder in full_text, f"Missing placeholder: {placeholder}"
