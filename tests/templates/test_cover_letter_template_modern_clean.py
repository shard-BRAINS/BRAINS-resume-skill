"""Validate the modern-clean cover-letter template is ATS-safe."""
from pathlib import Path

from scripts.validators.ats_check import ats_check

TEMPLATE = Path(__file__).parent.parent.parent / "templates" / "cover-letter" / "modern-clean.docx"


def test_template_exists():
    assert TEMPLATE.exists(), f"Template missing: {TEMPLATE}"


def test_template_passes_ats_check():
    result = ats_check(TEMPLATE)
    assert result.passed, f"Failures: {[f.code for f in result.failures]}"


def test_template_uses_informal_salutation_and_signoff():
    """Modern-clean uses 'Hi' and 'Best,' instead of 'Dear' and 'Sincerely,'."""
    from docx import Document
    doc = Document(str(TEMPLATE))
    flat = "\n".join(p.text for p in doc.paragraphs)
    assert "Hi {{SALUTATION}}" in flat
    assert "Best," in flat
    assert "Dear" not in flat
    assert "Sincerely" not in flat


def test_template_has_no_recipient_block():
    """Modern-clean omits the formal recipient address block."""
    from docx import Document
    doc = Document(str(TEMPLATE))
    flat = "\n".join(p.text for p in doc.paragraphs)
    assert "{{RECIPIENT_NAME}}" not in flat
    assert "{{RECIPIENT_TITLE}}" not in flat
    assert "{{RECIPIENT_COMPANY}}" not in flat
