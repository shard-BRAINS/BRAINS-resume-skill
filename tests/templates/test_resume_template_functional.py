"""Validate the functional resume template is ATS-safe."""
from pathlib import Path

from scripts.validators.ats_check import ats_check

TEMPLATE = Path(__file__).parent.parent.parent / "templates" / "resume" / "functional.docx"


def test_template_exists():
    assert TEMPLATE.exists(), f"Template missing: {TEMPLATE}"


def test_template_passes_ats_check():
    result = ats_check(TEMPLATE)
    assert result.passed, f"Failures: {[f.code for f in result.failures]}"


def test_template_contains_skills_heading_before_experience_heading():
    """Functional layout: skills come first, then experience."""
    from docx import Document
    doc = Document(str(TEMPLATE))
    text_in_order = [p.text for p in doc.paragraphs]
    flat = "\n".join(text_in_order)
    skills_pos = flat.find("Skills and Achievements")
    experience_pos = flat.find("Experience")
    assert 0 <= skills_pos < experience_pos


def test_template_uses_word_heading_styles():
    from docx import Document
    doc = Document(str(TEMPLATE))
    heading_paragraphs = [p for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert len(heading_paragraphs) >= 4  # Summary, Skills, Experience, Education
