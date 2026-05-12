"""Validate the executive resume template is ATS-safe."""
from pathlib import Path

from scripts.validators.ats_check import ats_check

TEMPLATE = Path(__file__).parent.parent.parent / "templates" / "resume" / "executive.docx"


def test_template_exists():
    assert TEMPLATE.exists(), f"Template missing: {TEMPLATE}"


def test_template_passes_ats_check():
    result = ats_check(TEMPLATE)
    assert result.passed, f"Failures: {[f.code for f in result.failures]}"


def test_template_uses_executive_section_labels():
    """Executive template uses 'Executive Summary' and 'Career Highlights'
    instead of the standard 'Summary' / 'Skills'."""
    from docx import Document
    doc = Document(str(TEMPLATE))
    flat = "\n".join(p.text for p in doc.paragraphs)
    assert "Executive Summary" in flat
    assert "Career Highlights" in flat


def test_template_name_heading_is_24pt():
    """Executive template uses 24pt name heading (2pt larger than chronological)."""
    from docx import Document
    doc = Document(str(TEMPLATE))
    name_paragraph = doc.paragraphs[0]
    assert name_paragraph.runs
    assert name_paragraph.runs[0].font.size.pt == 24
