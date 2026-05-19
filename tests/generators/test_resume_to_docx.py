"""Tests for the resume DOCX generator."""
from pathlib import Path
import pytest

from scripts.generators.resume_to_docx import render_resume_docx
from scripts.outputs.tagging import ArtifactMeta, read_artifact_meta
from scripts.validators.ats_check import ats_check


SAMPLE_RESUME_DATA = {
    "candidate_name": "Alex Test",
    "candidate_contact_line": "alex.test@example.invalid | Sample City",
    "summary": "Senior engineer with eight years of experience.",
    "skills": "Python, distributed systems, observability, mentoring",
    "experience": (
        "Senior Engineer, Example Corp 2020 - Present\n"
        "Built ingestion pipeline processing 50M events daily.\n\n"
        "Software Engineer, Sample Industries 2017 - 2019\n"
        "Backend services for an e-commerce platform."
    ),
    "education": "BSc Computer Science, Example University, 2015",
}


def test_render_writes_docx(tmp_path):
    out = tmp_path / "resume.docx"
    render_resume_docx(SAMPLE_RESUME_DATA, out)
    assert out.exists()
    assert out.stat().st_size > 0


def test_rendered_docx_contains_provided_content(tmp_path):
    out = tmp_path / "resume.docx"
    render_resume_docx(SAMPLE_RESUME_DATA, out)
    from docx import Document
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Alex Test" in text
    assert "Senior engineer" in text
    assert "Example Corp" in text
    assert "BSc Computer Science" in text


def test_rendered_docx_has_no_unfilled_placeholders(tmp_path):
    out = tmp_path / "resume.docx"
    render_resume_docx(SAMPLE_RESUME_DATA, out)
    from docx import Document
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "{{" not in text, "Template placeholder leaked into rendered output"


def test_rendered_docx_passes_ats_check(tmp_path):
    out = tmp_path / "resume.docx"
    render_resume_docx(SAMPLE_RESUME_DATA, out)
    result = ats_check(out)
    assert result.passed, f"Failures: {[f.code for f in result.failures]}"


def test_renders_to_user_chosen_directory(tmp_path):
    subdir = tmp_path / "myoutput"
    out = subdir / "resume.docx"
    render_resume_docx(SAMPLE_RESUME_DATA, out)
    assert out.exists()
    assert subdir.exists()


def test_render_accepts_template_argument(tmp_path):
    out = tmp_path / "resume.docx"
    render_resume_docx(SAMPLE_RESUME_DATA, out, template="chronological")
    assert out.exists()


def test_render_unknown_template_raises_value_error(tmp_path):
    out = tmp_path / "resume.docx"
    with pytest.raises(ValueError) as exc_info:
        render_resume_docx(SAMPLE_RESUME_DATA, out, template="nonexistent")
    assert "nonexistent" in str(exc_info.value)
    assert "chronological" in str(exc_info.value)  # message lists valid templates


def test_render_default_template_unchanged_behaviour(tmp_path):
    """Default behaviour matches v1.0.x — calling without template= produces
    the chronological-template output."""
    out_default = tmp_path / "default.docx"
    out_explicit = tmp_path / "explicit.docx"
    render_resume_docx(SAMPLE_RESUME_DATA, out_default)
    render_resume_docx(SAMPLE_RESUME_DATA, out_explicit, template="chronological")
    # Both produced output; both contain the candidate name.
    from docx import Document
    text_default = "\n".join(p.text for p in Document(str(out_default)).paragraphs)
    text_explicit = "\n".join(p.text for p in Document(str(out_explicit)).paragraphs)
    assert "Alex Test" in text_default
    assert "Alex Test" in text_explicit


def test_render_resume_docx_without_meta_still_works(tmp_path):
    """Backward compatibility: calling without artifact_meta= works as before."""
    out = tmp_path / "r.docx"
    render_resume_docx(
        data={"candidate_name": "M Gell"},
        out_path=out,
        template="chronological",
    )
    assert out.exists()
    assert read_artifact_meta(out) is None  # no custom properties


def test_render_resume_docx_with_meta_embeds_properties(tmp_path):
    """With artifact_meta provided, custom properties are embedded in the DOCX."""
    out = tmp_path / "r.docx"
    meta = ArtifactMeta(
        artifact_uid="KX7M9Q",
        artifact_kind="resume",
        jd_id=42,
        parent_uid=None,
        created_at="2026-05-19T10:00:00Z",
        skill_version="1.5.0",
    )
    render_resume_docx(
        data={"candidate_name": "M Gell"},
        out_path=out,
        template="chronological",
        artifact_meta=meta,
    )
    loaded = read_artifact_meta(out)
    assert loaded == meta
