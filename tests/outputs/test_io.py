"""Tests for scripts/outputs/io.py — orchestration layer."""
from pathlib import Path

import pytest

from scripts.outputs.io import (
    get_outputs_root,
    ProfileNameMissingError,
    OutputsDirNotWritableError,
    UIDCollisionError,
)


@pytest.fixture
def isolated(monkeypatch, tmp_path):
    """Each test gets isolated DB, profile, and outputs root."""
    monkeypatch.setenv("BRAINS_TRACKER_DB_PATH", str(tmp_path / "test.db"))
    monkeypatch.setenv("BRAINS_TRACKER_PROFILE_PATH", str(tmp_path / "profile.json"))
    monkeypatch.setenv("BRAINS_OUTPUTS_DIR", str(tmp_path / "outputs"))
    return tmp_path


def test_get_outputs_root_honours_env_var(isolated):
    assert get_outputs_root() == isolated / "outputs"


def test_get_outputs_root_defaults_when_unset(monkeypatch):
    monkeypatch.delenv("BRAINS_OUTPUTS_DIR", raising=False)
    assert get_outputs_root() == Path.home() / ".brains-resume" / "outputs"


def test_exceptions_are_distinct():
    # Sanity: the three exception classes exist and aren't the same.
    assert ProfileNameMissingError is not OutputsDirNotWritableError
    assert ProfileNameMissingError is not UIDCollisionError
    assert OutputsDirNotWritableError is not UIDCollisionError


from datetime import date

from scripts.outputs.io import ensure_jd_folder
from scripts.tracker.add import add_jd
from scripts.tracker.candidates import create_candidate, set_active_candidate
from scripts.tracker.db import open_db


def _ensure_active_candidate():
    """add_jd requires an active candidate; provide one for JD-only tests."""
    cid = create_candidate("Matthew", "Gell", [], None, None)
    set_active_candidate(cid)
    return cid


def test_ensure_jd_folder_creates_dir(isolated):
    _ensure_active_candidate()
    jd_id = add_jd(
        source="manual", source_ref=None,
        company="Acme Corp", role_title="Senior Data Engineer",
        raw_text="...", analyzer_findings={},
        focus_areas_required=[], focus_areas_nice=[],
    )
    folder = ensure_jd_folder(jd_id)
    assert folder.exists()
    assert folder.is_dir()
    # Folder name is YYYY-MM-DD_<anchor>_<role>
    assert "Acme-Corp" in folder.name
    assert "Senior-Data-Engineer" in folder.name


def test_ensure_jd_folder_writes_folder_path_to_jd_row(isolated):
    _ensure_active_candidate()
    jd_id = add_jd("manual", None, "Acme", "Eng", "...", {}, [], [])
    folder = ensure_jd_folder(jd_id)
    conn = open_db()
    try:
        row = conn.execute(
            "SELECT folder_path FROM jds WHERE id=?", (jd_id,),
        ).fetchone()
    finally:
        conn.close()
    assert row[0] == str(folder)


def test_ensure_jd_folder_idempotent(isolated):
    _ensure_active_candidate()
    jd_id = add_jd("manual", None, "Acme", "Eng", "...", {}, [], [])
    first = ensure_jd_folder(jd_id)
    second = ensure_jd_folder(jd_id)
    assert first == second


def test_ensure_jd_folder_handles_collision(isolated):
    """Two JDs added on the same day for the same company/role produce
    distinct folders via _v2 suffix."""
    _ensure_active_candidate()
    jd_a = add_jd("manual", None, "Acme", "Engineer", "...", {}, [], [])
    folder_a = ensure_jd_folder(jd_a)
    jd_b = add_jd("manual", None, "Acme", "Engineer", "...", {}, [], [])
    folder_b = ensure_jd_folder(jd_b)
    assert folder_a != folder_b
    assert folder_b.name.endswith("_v2")


from datetime import date as _date

from scripts.outputs.io import make_artifact_path
from scripts.outputs.tagging import ArtifactMeta
from scripts.tracker.add import NoActiveCandidateError
from scripts.tracker.candidates import create_candidate, set_active_candidate


def _activate_candidate(first="Matthew", last="Gell"):
    """Create a candidate and mark it active; return its id."""
    cid = create_candidate(first, last, [], None, None)
    set_active_candidate(cid)
    return cid


def test_make_artifact_path_resume(isolated):
    _activate_candidate()
    jd_id = add_jd("manual", None, "Acme", "Engineer", "...", {}, [], [])
    path, meta = make_artifact_path(jd_id, "resume")
    assert path.parent.exists()
    assert path.name.startswith("Matthew_Gell_resume_")
    assert path.suffix == ".docx"
    assert isinstance(meta, ArtifactMeta)
    assert meta.artifact_uid in path.name
    assert meta.artifact_kind == "resume"
    assert meta.jd_id == jd_id
    assert meta.parent_uid is None


def test_make_artifact_path_cover_letter(isolated):
    _activate_candidate()
    jd_id = add_jd("manual", None, "Acme", "Engineer", "...", {}, [], [])
    path, meta = make_artifact_path(jd_id, "cover-letter")
    assert path.name.startswith("Matthew_Gell_cover-letter_")


def test_make_artifact_path_with_parent_uid(isolated):
    _activate_candidate()
    jd_id = add_jd("manual", None, "Acme", "Engineer", "...", {}, [], [])
    path, meta = make_artifact_path(jd_id, "resume", parent_uid="ABCDEF")
    assert meta.parent_uid == "ABCDEF"


def test_make_artifact_path_raises_when_no_active_candidate(isolated):
    # A JD exists (created against its own candidate) but no active candidate
    # is set and no override is passed.
    cid = create_candidate("Matthew", "Gell", [], None, None)
    jd_id = add_jd(
        "manual", None, "Acme", "Engineer", "...", {}, [], [],
        candidate_id=cid,
    )
    with pytest.raises(NoActiveCandidateError):
        make_artifact_path(jd_id, "resume")


def test_make_artifact_path_two_calls_yield_distinct_uids(isolated):
    _activate_candidate()
    jd_id = add_jd("manual", None, "Acme", "Engineer", "...", {}, [], [])
    _, meta_a = make_artifact_path(jd_id, "resume")
    _, meta_b = make_artifact_path(jd_id, "resume")
    assert meta_a.artifact_uid != meta_b.artifact_uid


from docx import Document

from scripts.outputs.io import finalize_docx, read_artifact_uid
from scripts.outputs.tagging import read_artifact_meta


def test_finalize_docx_writes_custom_properties(isolated):
    _activate_candidate()
    jd_id = add_jd("manual", None, "Acme", "Engineer", "...", {}, [], [])
    path, meta = make_artifact_path(jd_id, "resume")
    # Generator writes the file:
    doc = Document()
    doc.add_paragraph("body")
    doc.save(str(path))
    # Skill finalizes:
    finalize_docx(path, meta)
    loaded = read_artifact_meta(path)
    assert loaded == meta


def test_finalize_docx_raises_when_file_missing(isolated):
    _activate_candidate()
    jd_id = add_jd("manual", None, "Acme", "Engineer", "...", {}, [], [])
    path, meta = make_artifact_path(jd_id, "resume")
    # Don't actually create the file.
    with pytest.raises(FileNotFoundError):
        finalize_docx(path, meta)


def test_read_artifact_uid_returns_uid_when_present(isolated):
    _activate_candidate()
    jd_id = add_jd("manual", None, "Acme", "Engineer", "...", {}, [], [])
    path, meta = make_artifact_path(jd_id, "resume")
    doc = Document()
    doc.save(str(path))
    finalize_docx(path, meta)
    assert read_artifact_uid(path) == meta.artifact_uid


def test_read_artifact_uid_returns_none_for_untagged(isolated, tmp_path):
    untagged = tmp_path / "u.docx"
    Document().save(str(untagged))
    assert read_artifact_uid(untagged) is None


def test_make_artifact_path_for_candidate_uses_candidate_name_in_filename(isolated):
    """for_candidate name supersedes the active candidate for filename construction."""
    _activate_candidate("Matthew", "Gell")
    jd_id = add_jd("manual", None, "Acme", "Sales Assistant", "...", {}, [], [])

    path, meta = make_artifact_path(
        jd_id, "resume", parent_uid=None,
        for_candidate="Mathilda Gell",
    )
    assert "Mathilda" in path.name
    assert "Gell" in path.name
    assert "Matthew" not in path.name
    assert meta.for_candidate == "Mathilda Gell"


def test_make_artifact_path_no_override_uses_active_candidate_name(isolated):
    """No override: the active candidate's name is used."""
    _activate_candidate("Matthew", "Gell")
    jd_id = add_jd("manual", None, "Acme", "Senior Eng", "...", {}, [], [])

    path, meta = make_artifact_path(jd_id, "resume", parent_uid=None)
    assert "Matthew" in path.name
    assert "Gell" in path.name
    assert meta.for_candidate is None


def test_make_artifact_path_single_token_candidate_name(isolated):
    """Single-token candidate names (Cher, Madonna) still produce a valid filename."""
    _activate_candidate("Matthew", "Gell")
    jd_id = add_jd("manual", None, "Acme", "Eng", "...", {}, [], [])

    path, meta = make_artifact_path(
        jd_id, "resume", parent_uid=None, for_candidate="Cher",
    )
    assert "Cher" in path.name
    assert meta.for_candidate == "Cher"


def test_make_artifact_path_uses_active_candidate_name(isolated):
    from scripts.outputs.io import make_artifact_path
    from scripts.tracker.add import add_jd
    from scripts.tracker.candidates import create_candidate, set_active_candidate

    cid = create_candidate("Mathilda", "Gell", [], None, None)
    set_active_candidate(cid)
    jd_id = add_jd("manual", None, "Acme", "Eng", "...", {}, [], [])

    path, meta = make_artifact_path(jd_id, "resume")
    assert "Mathilda" in path.name
    assert "Gell" in path.name
    assert meta.candidate_id == cid


def test_make_artifact_path_explicit_candidate_id_override(isolated):
    from scripts.outputs.io import make_artifact_path
    from scripts.tracker.add import add_jd
    from scripts.tracker.candidates import create_candidate, set_active_candidate

    matthew = create_candidate("Matthew", "Gell", [], None, None)
    mathilda = create_candidate("Mathilda", "Gell", [], None, None)
    set_active_candidate(matthew)
    jd_id = add_jd("manual", None, "Acme", "Eng", "...", {}, [], [])

    path, meta = make_artifact_path(jd_id, "resume", candidate_id=mathilda)
    assert "Mathilda" in path.name
    assert meta.candidate_id == mathilda


def test_make_artifact_path_for_candidate_alias_resolves_to_candidate_id(isolated):
    """v1.6 callers passing for_candidate still work; we resolve it to candidate_id."""
    from scripts.outputs.io import make_artifact_path
    from scripts.tracker.add import add_jd
    from scripts.tracker.candidates import create_candidate, set_active_candidate

    cid = create_candidate("Matthew", "Gell", [], None, None)
    set_active_candidate(cid)
    jd_id = add_jd("manual", None, "Acme", "Eng", "...", {}, [], [])

    path, meta = make_artifact_path(
        jd_id, "resume", for_candidate="Mathilda Gell",
    )
    assert "Mathilda" in path.name
    assert meta.for_candidate == "Mathilda Gell"
    assert meta.candidate_id is not None
    assert meta.candidate_id != cid  # not Matthew
