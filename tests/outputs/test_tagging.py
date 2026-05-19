"""Tests for scripts/outputs/tagging.py — DOCX custom property read/write."""
from pathlib import Path

import pytest
from docx import Document

from scripts.outputs.tagging import ArtifactMeta, write_artifact_meta, read_artifact_meta, _read_custom_properties_raw


@pytest.fixture
def minimal_docx(tmp_path):
    """A bare DOCX with one paragraph."""
    path = tmp_path / "minimal.docx"
    doc = Document()
    doc.add_paragraph("Hello world.")
    doc.save(str(path))
    return path


def _meta_for_test():
    return ArtifactMeta(
        artifact_uid="KX7M9Q",
        artifact_kind="resume",
        jd_id=47,
        parent_uid="PT4N2B",
        created_at="2026-05-19T10:33:02Z",
        skill_version="1.5.0",
    )


def test_write_artifact_meta_creates_custom_properties(minimal_docx):
    write_artifact_meta(minimal_docx, _meta_for_test())
    props = _read_custom_properties_raw(minimal_docx)
    assert props["BrainsArtifactId"] == "KX7M9Q"
    assert props["BrainsArtifactKind"] == "resume"
    assert props["BrainsJDId"] == "47"  # raw XML is a string
    assert props["BrainsParentId"] == "PT4N2B"
    assert props["BrainsCreatedAt"] == "2026-05-19T10:33:02Z"
    assert props["BrainsSkillVersion"] == "1.5.0"


def test_write_artifact_meta_with_no_parent(minimal_docx):
    meta = _meta_for_test()
    meta.parent_uid = None
    write_artifact_meta(minimal_docx, meta)
    props = _read_custom_properties_raw(minimal_docx)
    # Absent parent stored as empty string.
    assert props["BrainsParentId"] == ""


def test_write_artifact_meta_with_no_jd(minimal_docx):
    meta = _meta_for_test()
    meta.jd_id = None
    write_artifact_meta(minimal_docx, meta)
    props = _read_custom_properties_raw(minimal_docx)
    # Absent JD stored as 0.
    assert props["BrainsJDId"] == "0"


def test_write_artifact_meta_does_not_change_body(minimal_docx):
    original_text = Document(str(minimal_docx)).paragraphs[0].text
    write_artifact_meta(minimal_docx, _meta_for_test())
    after_text = Document(str(minimal_docx)).paragraphs[0].text
    assert original_text == after_text == "Hello world."


def test_write_artifact_meta_docx_remains_valid(minimal_docx):
    """python-docx must still be able to open the file after we add custom properties."""
    write_artifact_meta(minimal_docx, _meta_for_test())
    doc = Document(str(minimal_docx))  # should not raise
    assert len(doc.paragraphs) == 1


def test_write_artifact_meta_overwrites_existing(minimal_docx):
    """Calling write_artifact_meta twice updates rather than duplicating."""
    write_artifact_meta(minimal_docx, _meta_for_test())
    meta2 = _meta_for_test()
    meta2.artifact_uid = "DIFFRT"
    write_artifact_meta(minimal_docx, meta2)
    props = _read_custom_properties_raw(minimal_docx)
    assert props["BrainsArtifactId"] == "DIFFRT"


def test_read_artifact_meta_round_trip(minimal_docx):
    original = _meta_for_test()
    write_artifact_meta(minimal_docx, original)
    loaded = read_artifact_meta(minimal_docx)
    assert loaded == original


def test_read_artifact_meta_converts_zero_jd_to_none(minimal_docx):
    meta = _meta_for_test()
    meta.jd_id = None
    write_artifact_meta(minimal_docx, meta)
    loaded = read_artifact_meta(minimal_docx)
    assert loaded.jd_id is None


def test_read_artifact_meta_converts_empty_parent_to_none(minimal_docx):
    meta = _meta_for_test()
    meta.parent_uid = None
    write_artifact_meta(minimal_docx, meta)
    loaded = read_artifact_meta(minimal_docx)
    assert loaded.parent_uid is None


def test_read_artifact_meta_returns_none_on_untagged_docx(minimal_docx):
    # No write_artifact_meta call — fresh docx.
    assert read_artifact_meta(minimal_docx) is None


def test_read_artifact_meta_survives_rename(minimal_docx, tmp_path):
    write_artifact_meta(minimal_docx, _meta_for_test())
    renamed = tmp_path / "manually-renamed.docx"
    minimal_docx.rename(renamed)
    loaded = read_artifact_meta(renamed)
    assert loaded.artifact_uid == "KX7M9Q"


def test_artifact_meta_has_for_candidate_default_none():
    from scripts.outputs.tagging import ArtifactMeta
    m = ArtifactMeta(
        artifact_uid="ABC123", artifact_kind="resume",
        jd_id=None, parent_uid=None,
        created_at="2026-05-19T00:00:00Z", skill_version="1.5.0",
    )
    assert m.for_candidate is None


def test_for_candidate_roundtrips_through_docx(tmp_path):
    """Write a DOCX with for_candidate set, read it back."""
    from docx import Document
    from scripts.outputs.tagging import ArtifactMeta, write_artifact_meta, read_artifact_meta

    docx_path = tmp_path / "test.docx"
    Document().save(str(docx_path))

    meta = ArtifactMeta(
        artifact_uid="ABC123", artifact_kind="resume",
        jd_id=None, parent_uid=None,
        created_at="2026-05-19T00:00:00Z", skill_version="1.5.0",
        for_candidate="Mathilda Gell",
    )
    write_artifact_meta(docx_path, meta)
    roundtrip = read_artifact_meta(docx_path)
    assert roundtrip.for_candidate == "Mathilda Gell"


def test_for_candidate_absent_reads_as_none(tmp_path):
    """A DOCX written without for_candidate must read as None (forward-compat)."""
    from docx import Document
    from scripts.outputs.tagging import ArtifactMeta, write_artifact_meta, read_artifact_meta

    docx_path = tmp_path / "test.docx"
    Document().save(str(docx_path))

    meta = ArtifactMeta(
        artifact_uid="ABC123", artifact_kind="resume",
        jd_id=None, parent_uid=None,
        created_at="2026-05-19T00:00:00Z", skill_version="1.5.0",
        # for_candidate omitted
    )
    write_artifact_meta(docx_path, meta)
    roundtrip = read_artifact_meta(docx_path)
    assert roundtrip.for_candidate is None
