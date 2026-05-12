"""Generate the functional resume DOCX template.

Skills-led layout. Experience section is minimal — titles + dates only.
Single-column. No tables. Calibri 11pt body, Calibri Light 22pt name.
Margins 0.75 inch. Page size US Letter. ATS-safe.
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

TEMPLATE_PATH = Path(__file__).parent.parent.parent / "templates" / "resume" / "functional.docx"


def make_template() -> Path:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = name.add_run("{{CANDIDATE_NAME}}")
    run.font.name = "Calibri Light"
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)

    contact = doc.add_paragraph("{{CANDIDATE_CONTACT_LINE}}")
    contact.runs[0].font.size = Pt(10)
    contact.runs[0].font.color.rgb = RGBColor(0x40, 0x40, 0x40)

    doc.add_paragraph()

    doc.add_heading("Summary", level=1)
    doc.add_paragraph("{{SUMMARY}}")

    # Functional layout: skills FIRST, with skill-grouped achievements.
    doc.add_heading("Skills and Achievements", level=1)
    doc.add_paragraph("{{SKILLS}}")

    # Experience section is minimised — titles + dates only.
    doc.add_heading("Experience", level=1)
    doc.add_paragraph("{{EXPERIENCE}}")

    doc.add_heading("Education", level=1)
    doc.add_paragraph("{{EDUCATION}}")

    TEMPLATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(TEMPLATE_PATH))
    return TEMPLATE_PATH


if __name__ == "__main__":
    path = make_template()
    print(f"Wrote {path}")
