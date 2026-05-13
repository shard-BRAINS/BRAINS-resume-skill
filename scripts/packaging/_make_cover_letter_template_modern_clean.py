"""Generate the modern-clean cover-letter DOCX template.

Less formal layout. Smaller letterhead, less typographic weight, more
whitespace between paragraphs. Suited to tech/startup contexts where the
formal letterhead reads as stiff. Single-column. Calibri 11pt body. ATS-safe.
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor

TEMPLATE_PATH = Path(__file__).parent.parent.parent / "templates" / "cover-letter" / "modern-clean.docx"


def make_template() -> Path:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)  # more top space than formal-business
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Compact letterhead — name slightly larger but no full recipient block
    p = doc.add_paragraph()
    run = p.add_run("{{CANDIDATE_NAME}}")
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    doc.add_paragraph("{{CANDIDATE_CONTACT_LINE}}")
    doc.add_paragraph()
    doc.add_paragraph()  # extra whitespace

    doc.add_paragraph("{{LETTER_DATE}}")
    doc.add_paragraph()

    # Less-formal salutation — "Hi" instead of "Dear"
    doc.add_paragraph("Hi {{SALUTATION}},")
    doc.add_paragraph()

    doc.add_paragraph("{{HOOK_PARAGRAPH}}")
    doc.add_paragraph()
    doc.add_paragraph("{{FIT_PARAGRAPH}}")
    doc.add_paragraph()
    doc.add_paragraph("{{CLOSE_PARAGRAPH}}")
    doc.add_paragraph()

    # Less-formal sign-off
    doc.add_paragraph("Best,")
    doc.add_paragraph()
    doc.add_paragraph("{{CANDIDATE_NAME}}")

    TEMPLATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(TEMPLATE_PATH))
    return TEMPLATE_PATH


if __name__ == "__main__":
    path = make_template()
    print(f"Wrote {path}")
