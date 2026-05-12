"""Generate the ATS-safe cover-letter DOCX template.

Formal business-letter layout. Single column. Same font discipline as the
resume template (Calibri 11pt body, Heading 1 for any structural headers).
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor

TEMPLATE_PATH = Path(__file__).parent.parent.parent / "templates" / "cover-letter" / "formal-business.docx"


def make_template() -> Path:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    p = doc.add_paragraph()
    run = p.add_run("{{CANDIDATE_NAME}}")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    doc.add_paragraph("{{CANDIDATE_CONTACT_LINE}}")
    doc.add_paragraph()

    doc.add_paragraph("{{LETTER_DATE}}")
    doc.add_paragraph()

    doc.add_paragraph("{{RECIPIENT_NAME}}")
    doc.add_paragraph("{{RECIPIENT_TITLE}}")
    doc.add_paragraph("{{RECIPIENT_COMPANY}}")
    doc.add_paragraph()

    doc.add_paragraph("Dear {{SALUTATION}},")
    doc.add_paragraph()

    doc.add_paragraph("{{HOOK_PARAGRAPH}}")
    doc.add_paragraph("{{FIT_PARAGRAPH}}")
    doc.add_paragraph("{{CLOSE_PARAGRAPH}}")
    doc.add_paragraph()

    doc.add_paragraph("Sincerely,")
    doc.add_paragraph()
    doc.add_paragraph("{{CANDIDATE_NAME}}")

    TEMPLATE_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(TEMPLATE_PATH))
    return TEMPLATE_PATH


if __name__ == "__main__":
    path = make_template()
    print(f"Wrote {path}")
