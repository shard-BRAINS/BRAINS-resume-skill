"""Resume DOCX generator.

Takes structured resume content as input, fills the chronological
template's placeholders, writes a clean ATS-safe DOCX. UNBRANDED — this
is the user's professional document, not a BRAINS coaching artifact.
"""
from pathlib import Path
from typing import Union

from docx import Document

DEFAULT_TEMPLATE = "chronological"
TEMPLATES_DIR = Path(__file__).parent.parent.parent / "templates" / "resume"


PLACEHOLDER_MAP = {
    "candidate_name": "{{CANDIDATE_NAME}}",
    "candidate_contact_line": "{{CANDIDATE_CONTACT_LINE}}",
    "summary": "{{SUMMARY}}",
    "skills": "{{SKILLS}}",
    "experience": "{{EXPERIENCE}}",
    "education": "{{EDUCATION}}",
}


def _substitute_placeholder(paragraph, placeholder: str, value: str):
    """Replace a placeholder in a paragraph, preserving the first run's formatting."""
    if placeholder not in paragraph.text:
        return False
    if not paragraph.runs:
        return False
    first_run = paragraph.runs[0]
    full_text = paragraph.text.replace(placeholder, value)
    for run in paragraph.runs[1:]:
        run.text = ""
    first_run.text = full_text
    return True


def _expand_multiline_paragraph(doc: Document, paragraph, value: str):
    """If the value contains newlines, split into multiple paragraphs.

    Inserts new paragraphs after the original, preserving sequence.
    """
    if "\n" not in value:
        return
    lines = value.split("\n")
    first_line = lines[0]
    rest = lines[1:]
    paragraph.runs[0].text = first_line
    p_element = paragraph._element
    parent = p_element.getparent()
    insertion_index = list(parent).index(p_element) + 1
    for line in rest:
        new_p = doc.add_paragraph(line)
        new_p_element = new_p._element
        parent.remove(new_p_element)
        parent.insert(insertion_index, new_p_element)
        insertion_index += 1


def render_resume_docx(data: dict, out_path: Union[str, Path]) -> Path:
    """Render a structured resume dict into the chronological DOCX template.

    Required keys in data: candidate_name, candidate_contact_line, summary,
    skills, experience, education. Missing keys default to an empty string.
    """
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    template_path = TEMPLATES_DIR / f"{DEFAULT_TEMPLATE}.docx"
    if not template_path.exists():
        raise FileNotFoundError(f"Resume template not found: {template_path}")

    doc = Document(str(template_path))

    multiline_jobs = []
    for paragraph in list(doc.paragraphs):
        for key, placeholder in PLACEHOLDER_MAP.items():
            if placeholder in paragraph.text:
                value = data.get(key, "") or ""
                if "\n" in value:
                    multiline_jobs.append((paragraph, value))
                else:
                    _substitute_placeholder(paragraph, placeholder, value)

    for paragraph, value in multiline_jobs:
        _expand_multiline_paragraph(doc, paragraph, value)

    doc.save(str(out_path))
    return out_path
