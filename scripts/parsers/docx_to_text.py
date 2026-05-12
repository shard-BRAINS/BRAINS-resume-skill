"""DOCX resume parser.

Extracts text and detects section structure from a resume DOCX using python-docx.
Returns a dict with raw text, detected sections, and paragraph count.

Heading detection uses three signals in priority order:
  1. Real Word Heading 1/2/3/Title styles (strongest signal)
  2. A bold run + larger-than-body font size (>= 13 pt by default)
  3. Short paragraph (<= 40 chars) whose lowercased text matches the
     SECTION_KEYWORDS vocabulary

Any of the three is sufficient to mark a paragraph as a section heading.
"""
from pathlib import Path
from typing import Union

from docx import Document
from docx.shared import Pt

SECTION_KEYWORDS = {
    "summary", "profile", "objective", "about",
    "experience", "employment", "work history", "professional experience",
    "education", "academic", "qualifications",
    "skills", "technical skills", "competencies",
    "projects", "publications", "certifications", "languages",
    "achievements", "awards", "interests", "volunteer",
}

BODY_FONT_PT_THRESHOLD = 13  # font sizes >= this on bold runs count as heading


def _paragraph_is_styled_heading(paragraph) -> bool:
    """True if the paragraph uses a Word Heading or Title style."""
    style_name = (paragraph.style.name or "").lower()
    return style_name.startswith("heading") or style_name == "title"


def _paragraph_has_bold_large_run(paragraph) -> bool:
    """True if the paragraph contains a bold run at >= BODY_FONT_PT_THRESHOLD points.

    Falls back to False if font size is not explicitly set (cannot infer).
    """
    for run in paragraph.runs:
        if not run.bold:
            continue
        size = run.font.size
        if size is None:
            continue
        if size.pt >= BODY_FONT_PT_THRESHOLD:
            return True
    return False


def _matches_keyword_vocab(text: str) -> bool:
    """True if a short paragraph matches the section-keyword vocabulary."""
    stripped = text.strip()
    if len(stripped) > 40 or not stripped:
        return False
    lower = stripped.lower().rstrip(":")
    return lower in SECTION_KEYWORDS


def _canonicalise_heading(text: str) -> str:
    """Canonicalise heading text for use as a section key."""
    return text.strip().lower().rstrip(":")


def parse_docx_resume(path: Union[str, Path]) -> dict:
    """Parse a resume DOCX and return structured content."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"DOCX not found: {path}")

    doc = Document(str(path))
    raw_text_parts: list = []
    sections: dict = {}
    current_heading = "_preamble"
    current_lines: list = []

    for paragraph in doc.paragraphs:
        text = paragraph.text
        raw_text_parts.append(text)

        is_styled_heading = _paragraph_is_styled_heading(paragraph)
        is_bold_large = _paragraph_has_bold_large_run(paragraph)
        is_keyword_match = _matches_keyword_vocab(text)

        if text.strip() and (is_styled_heading or is_bold_large or is_keyword_match):
            if current_lines:
                sections[current_heading] = "\n".join(current_lines).strip()
            current_heading = _canonicalise_heading(text)
            current_lines = []
        else:
            current_lines.append(text)

    if current_lines:
        sections[current_heading] = "\n".join(current_lines).strip()

    return {
        "raw_text": "\n".join(raw_text_parts),
        "sections": sections,
        "paragraph_count": len(doc.paragraphs),
    }
