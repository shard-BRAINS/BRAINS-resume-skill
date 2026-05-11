"""Generate ATS-check fixtures: a clean DOCX, one with a table, one with a header."""
from pathlib import Path
from docx import Document
from docx.shared import Pt

here = Path(__file__).parent

# ats_clean.docx — should PASS all checks
doc = Document()
doc.add_heading("Alex Test", level=0)
doc.add_paragraph("alex.test@example.invalid")
doc.add_heading("Summary", level=1)
doc.add_paragraph("Software engineer.")
doc.add_heading("Experience", level=1)
doc.add_paragraph("Senior Engineer, Example Corp, 2020 — Present.")
doc.save(str(here / "ats_clean.docx"))

# ats_with_table.docx — should FAIL the no-tables check
doc = Document()
doc.add_heading("Alex Test", level=0)
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = "Skill"
table.cell(0, 1).text = "Years"
table.cell(1, 0).text = "Python"
table.cell(1, 1).text = "8"
doc.save(str(here / "ats_with_table.docx"))

# ats_with_header.docx — should FAIL the critical-info-in-header check
doc = Document()
section = doc.sections[0]
header = section.header
header.paragraphs[0].text = "Alex Test · alex.test@example.invalid · +0 0000 000000"
doc.add_heading("Summary", level=1)
doc.add_paragraph("Software engineer.")
doc.save(str(here / "ats_with_header.docx"))

print("Wrote three ATS fixtures.")
