"""Generate a synthetic resume DOCX for testing. Run once; commit the resulting DOCX.
NO REAL PII — all content is fictional and innocuous."""
from pathlib import Path
from docx import Document

fixture_path = Path(__file__).parent / "synthetic_resume_basic.docx"

doc = Document()
doc.add_heading("Alex Test", level=0)
doc.add_paragraph("alex.test@example.invalid · +0 0000 000000")
doc.add_heading("Summary", level=1)
doc.add_paragraph("Software engineer with eight years of experience building data systems.")
doc.add_heading("Experience", level=1)
doc.add_paragraph("Senior Engineer, Example Corp · 2020 — Present")
doc.add_paragraph("Built ingestion pipeline processing 50M events daily.")
doc.add_heading("Education", level=1)
doc.add_paragraph("BSc Computer Science, Example University, 2015")
doc.save(str(fixture_path))
print(f"Wrote {fixture_path}")
